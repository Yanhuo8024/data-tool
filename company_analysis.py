import logging

logger = logging.getLogger(__name__)
#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
company_analysis.py - 公司全方位分析
输入公司名（中文/代码）→ 聚合多维度真实数据 → 输出结构化报告 + 生成 Word 文档

覆盖章节：公司概况 / 财务分析 / 市场表现 / 最新新闻与舆情 / 竞争对手对比 / SWOT / 风险提示

数据源：akshare（A股/港股）。复用 financial_statement / valuation 现有函数。
"""
import re
import io
import base64
from pathlib import Path
from datetime import datetime

try:
    from core_config import OUT_DIR  # 统一入口，避免多处定义分叉
except ImportError:
    OUT_DIR = Path(__file__).parent / "output"  # core_config 不可用时兜底
OUT_DIR.mkdir(exist_ok=True)

def _resolve_cn_name(code: str):
    """股票代码 → 中文名称。依次尝试东财全市场快照表 / 新浪日线，失败返回空。"""
    try:
        import akshare as ak
        raw = code[2:] if str(code).startswith('HK') else code
        if str(code).startswith('HK'):
            df = ak.stock_zh_a_spot_em(symbol='港股')
            if df is not None and not df.empty:
                hit = df[df['代码'].astype(str).str.zfill(5).str.endswith(raw)]
                if not hit.empty:
                    nm = str(hit.iloc[0].get('名称', ''))
                    if nm:
                        return nm
            df = ak.stock_hk_spot_em()
            if df is not None and not df.empty:
                hit = df[df['代码'].astype(str).str.zfill(5).str.endswith(raw)]
                if not hit.empty:
                    c = df.columns
                    nm = str(hit.iloc[0].get('名称', '')) or str(hit.iloc[0].get(c[1], ''))
                    if nm:
                        return nm
            return ""
        df = ak.stock_zh_a_spot_em()
        if df is not None and not df.empty:
            hit = df[df['代码'].astype(str) == raw]
            if not hit.empty:
                nm = str(hit.iloc[0].get('名称', ''))
                if nm:
                    return nm
        sina = ('sh' if raw.startswith('6') else 'sz') + raw
        kdf = ak.stock_zh_a_daily(symbol=sina, adjust="")
        if kdf is not None and hasattr(kdf, 'attrs') and kdf.attrs.get('name'):
            return kdf.attrs['name']
    except Exception as _e:
        logger.debug("company_analysis.py:57 swallowed: %s", _e)
    return ""

import sys
sys.path.insert(0, str(Path(__file__).parent))

from financial_statement import fetch_financial_data, calculate_ratios, _resolve_stock_code, _resolve_stock_name
from valuation import _snapshot, dcf_valuation, peer_comparison

from datetime import timedelta

# 常见公司名→代码（覆盖 AI 助手主力标的；扩展时同步 ai_assistant._STOCK_ALIASES）
STOCK_ALIASES = {
    '贵州茅台': '600519', '茅台': '600519', '五粮液': '000858', '泸州老窖': '000568',
    '山西汾酒': '600809', '宁德时代': '300750', '隆基绿能': '601012', '中国平安': '601318',
    '招商银行': '600036', '腾讯': 'HK0700', '阿里巴巴': 'HK9988', '美团': 'HK03690',
    '小米': 'HK01810', '比亚迪': '002594', '中兴通讯': '000063', '兆易创新': '603986',
    '韦尔股份': '603501', '寒武纪': '688256', '海光信息': '688041', '中芯国际': '688981',
    '恒瑞医药': '600276', '海康威视': '002415', '东方财富': '300059', '科大讯飞': '002230',
    '工业富联': '601138', '紫光国微': '002049', '三安光电': '600703', '北方华创': '002371',
    '立讯精密': '002475', '牧原股份': '002714', '万科': '000002', '京东方': '000725',
    '星宇股份': '601799', '星宇': '601799',
}

# 巨潮证券搜索缓存（进程内）
_CNINFO_TOPSEARCH_CACHE = {}

def _cninfo_top_search(keyword: str) -> list:
    """巨潮证券代码搜索：公司名/拼音/代码 → 证券信息列表（别名表之外的动态解析来源）"""
    kw = (keyword or "").strip()
    if not kw:
        return []
    if kw in _CNINFO_TOPSEARCH_CACHE:
        return _CNINFO_TOPSEARCH_CACHE[kw]
    out = []
    try:
        import requests
        r = requests.post(
            "http://www.cninfo.com.cn/new/information/topSearch/query",
            data={"keyWord": kw, "maxNum": 10},
            headers=_CNINFO_HEADERS, timeout=8)
        arr = r.json()
        if isinstance(arr, list):
            out = [it for it in arr if isinstance(it, dict)]
    except Exception as _e:
        logger.debug("cninfo topSearch swallowed: %s", _e)
    _CNINFO_TOPSEARCH_CACHE[kw] = out
    return out

def _resolve_via_cninfo(name: str):
    """巨潮动态解析：公司名 → (code, market) 或 (None, None)

    只接受**精确简称**匹配且分类为 A股/港股 —— 长句切出的行业词/普通词
    （如"新能源汽车"）若允许模糊匹配会误命中"宝新能源"等证券，把行业
    分析消息误路由到个股取数。
    """
    arr = _cninfo_top_search(name)
    if not arr:
        return None, None
    nm = (name or "").strip()
    for it in arr:
        if str(it.get("zwjc") or "").strip() != nm:
            continue
        if str(it.get("delisted") or "") == "true":
            continue
        code_i = str(it.get("code") or "").strip()
        cat = str(it.get("category") or "")
        if not code_i:
            continue
        if cat in ("港股", "港股通") or (code_i.isdigit() and len(code_i) == 5):
            return "HK" + code_i.lstrip('0').zfill(4), "hk"
        if cat == "A股" or (code_i.isdigit() and len(code_i) == 6):
            return code_i, "a"
    return None, None

# 内置可比公司映射（同行业对比用）
PEERS_BY_CODE = {
    '600519': ['000858', '000568', '600809'], '000858': ['600519', '000568', '600809'],
    '600036': ['601398', '601939', '000001'], '601318': ['601398', '601939', '600036'],
    'HK0700': ['HK9988', 'HK03690', 'HK01810'], '300750': ['002594', '601127'],
    '600941': ['600050', '601728'], '002594': ['300750', '601127'],
}

def resolve_ticker(company_name: str):
    """公司名/代码 → 统一标识（HK0700 / 6位A股）。返回 (code, market) 或 (None,None)

    解析顺序：显式代码 → 别名表 → 巨潮证券搜索动态解析（覆盖任意 A股/港股公司名）。
    """
    name = (company_name or '').strip()
    if not name:
        return None, None
    # 直接是代码
    m = re.search(r'HK\s?(\d{4,5})', name, re.I)
    if m:
        return 'HK' + m.group(1).zfill(4), 'hk'
    m = re.search(r'\b(\d{6})\b', name)
    if m:
        c = m.group(1)
        return c, 'a'
    # 5位纯数字（港股代码写法，如 00700）→ 统一为 HK+4 位惯例（HK0700）
    if re.fullmatch(r'\d{5}', name):
        return 'HK' + name.lstrip('0').zfill(4), 'hk'
    # 公司名映射
    for k, code in STOCK_ALIASES.items():
        if k in name or name in k:
            return ('HK' + code[2:]) if str(code).startswith('HK') else code, ('hk' if str(code).startswith('HK') else 'a')
    # 别名表未命中 → 巨潮动态解析
    code, mkt = _resolve_via_cninfo(name)
    if code:
        return code, mkt
    return None, None

def code_to_name(code: str) -> str:
    """股票代码反转名称（优先取别名表最长匹配，保证中文名稳定显示；巨潮反查作为补充）"""
    for k, c in STOCK_ALIASES.items():
        if c == code:
            # 取最长名（避免 茅台/贵州茅台 取舍）
            return max([k for kk, cc in STOCK_ALIASES.items() if cc == code], key=len)
    # 别名表外：巨潮反查简称（东财接口不可达时的可靠来源）
    raw = str(code or '').strip()
    if raw:
        q = raw[2:] if raw.upper().startswith('HK') else raw
        for it in _cninfo_top_search(q):
            ci = str(it.get('code') or '').strip()
            if (ci == q or (q.isdigit() and ci.zfill(5) == q.zfill(5))) and it.get('zwjc'):
                return str(it['zwjc'])
    return code

_CNINFO_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0 Safari/537.36",
    "Referer": "http://www.cninfo.com.cn/",
    "Accept-Language": "zh-CN,zh;q=0.9",
}

def _cninfo_intro(code: str) -> dict:
    """巨潮资料卡：A股简称/行业/主营/上市日期/总股本。

    东财个股信息接口（stock_individual_info_em）在本机网络环境不可达时的兜底来源。
    """
    out = {}
    try:
        import requests
        r = requests.get(
            "http://www.cninfo.com.cn/data20/companyOverview/getCompanyIntroduction",
            params={"scode": str(code)}, headers=_CNINFO_HEADERS, timeout=8)
        recs = ((r.json().get("data") or {}).get("records")) or []
        bi = (recs[0].get("basicInformation") or [{}])[0] if recs else {}
        if bi:
            out["name"] = bi.get("ASECNAME") or ""
            out["industry"] = bi.get("F032V") or ""
            out["desc"] = bi.get("F015V") or bi.get("F016V") or ""
            out["list_time"] = bi.get("F006D") or bi.get("F010D") or ""
            ts = bi.get("F007N")  # 注册资本/总股本（万股）
            if ts:
                try:
                    out["total_share_yi"] = round(float(ts) / 1e4, 2)
                except (TypeError, ValueError):
                    pass
    except Exception:
        pass
    return out

def _em_f10_profile(code: str) -> dict:
    """东财数据中心 F10 公司资料：行业 + 主营（A股/港股/北交所通用）。

    datacenter-web 域名在本机网络环境可达（业绩报表同域）。
    港股行业取 BELONG_INDUSTRY（东财港股行业），A股取 INDUSTRYCSRC1（证监会分类，
    与巨潮行业聚合表口径一致）。
    """
    out = {}
    c = str(code or '').strip()
    try:
        import requests
        if c.upper().startswith('HK'):
            secucode = c[2:].zfill(5) + '.HK'
            report = 'RPT_HKF10_INFO_ORGPROFILE'
        elif c.isdigit() and len(c) == 6:
            suffix = '.BJ' if c.startswith(('4', '8')) else ('.SH' if c.startswith(('6', '9')) else '.SZ')
            secucode = c + suffix
            report = 'RPT_F10_BASIC_ORGINFO'
        else:
            return out
        h = {**_CNINFO_HEADERS, 'Referer': 'https://emweb.securities.eastmoney.com/'}
        r = requests.get(
            'https://datacenter-web.eastmoney.com/api/data/v1/get',
            params={'reportName': report, 'columns': 'ALL', 'pageSize': '1',
                    'filter': f'(SECUCODE="{secucode}")'},
            headers=h, timeout=8)
        rec = ((r.json().get('result') or {}).get('data') or [{}])[0]
        if not rec:
            return out
        if report == 'RPT_HKF10_INFO_ORGPROFILE':
            out['industry'] = str(rec.get('BELONG_INDUSTRY') or '').strip()
            it = str(rec.get('INDUSTRY_TYPE') or '').strip()
            if it:
                out['industry_detail'] = it.split('-')[-1]
            out['desc'] = str(rec.get('MAIN_BUSINESS') or '').strip()
            out['name'] = str(rec.get('SECURITY_NAME_ABBR') or '').strip()
        else:
            csrc = str(rec.get('INDUSTRYCSRC1') or '').strip()
            out['industry'] = csrc.split('-')[-1] if csrc else str(rec.get('EM2016') or '').split('-')[-1]
            em = str(rec.get('EM2016') or '').strip()
            if em:
                out['industry_detail'] = em.split('-')[-1]
            out['desc'] = str(rec.get('MAIN_BUSINESS') or rec.get('OPERATE_ADDRESS') or '').strip()
            out['name'] = str(rec.get('SECURITY_NAME_ABBR') or '').strip()
    except Exception:
        pass
    return out

def get_company_profile(code: str, market: str) -> dict:
    """公司概况：名称/行业/市值/上市时间/股本等（东财个股信息，失败则巨潮资料卡/名称兜底）"""
    out = {"name": code, "industry": "", "market_cap_yi": None, "list_time": "",
           "total_share_yi": None, "desc": ""}
    try:
        import akshare as ak
        q = ak.stock_individual_info_em(symbol=code)
        if isinstance(q, dict) and q:
            out["name"] = q.get("股票简称") or q.get("股票名称") or out["name"]
            out["industry"] = q.get("行业") or ""
            out["list_time"] = str(q.get("上市时间") or "")
            mc = q.get("总市值(元)")
            if mc:
                try:
                    out["market_cap_yi"] = round(float(mc) / 1e8, 1)
                except (ValueError, TypeError):
                    out["market_cap_yi"] = None
            ts = q.get("总股本(股)")
            if ts:
                try:
                    out["total_share_yi"] = round(float(ts) / 1e8, 2)
                except (ValueError, TypeError):
                    out["total_share_yi"] = None
    except Exception:
        pass
    # 巨潮资料卡兜底（补行业/简称/主营/上市日期/总股本）
    if (not out["industry"]) or (not out["name"]) or out["name"] == code \
            or out["total_share_yi"] is None:
        intro = _cninfo_intro(code)
        for k in ("industry", "desc", "list_time"):
            if not out.get(k) and intro.get(k):
                out[k] = intro[k]
        if (not out["name"] or out["name"] == code) and intro.get("name"):
            out["name"] = intro["name"]
        if out["total_share_yi"] is None and intro.get("total_share_yi"):
            out["total_share_yi"] = intro["total_share_yi"]
    # 东财数据中心 F10 兜底（A股/港股/北交所通用；巨潮不支持港股时的关键来源）
    if (not out["industry"]) or (not out["name"]) or out["name"] == code:
        f10 = _em_f10_profile(code)
        if not out["industry"] and f10.get("industry"):
            out["industry"] = f10["industry"]
        if not out.get("desc") and f10.get("desc"):
            out["desc"] = f10["desc"]
        if (not out["name"] or out["name"] == code) and f10.get("name"):
            out["name"] = f10["name"]
    # 名称兜底
    if not out["name"] or out["name"] == code:
        try:
            cn = _resolve_cn_name(code)
            if cn:
                out["name"] = cn
        except Exception:
            pass
    return out

def _amt_to_yi(v):
    """报表数值（'922.78亿' / '1.20万' / 原始元值）→ 亿元 float；无效返回 None"""
    try:
        s = str(v).replace(',', '').strip()
        if s.endswith('亿'):
            return round(float(s[:-1]), 2)
        if s.endswith('万'):
            return round(float(s[:-1]) / 1e4, 4)
        x = float(s)
        return round(x / 1e8, 2) if abs(x) >= 1e6 else round(x, 4)
    except (TypeError, ValueError):
        return None

def _extract_statements(data: dict) -> dict:
    """新浪三表 records → 最新期绝对值（亿元）+ 应收/存货同比 + 毛利率同比变化。

    供企业深析流水线回填财务表单（资产/负债/现金流/商誉等硬数据）。
    """
    out = {}
    dd = data.get("data") or {}

    def sheet(key):
        recs = dd.get(key) or []
        try:
            return sorted(recs, key=lambda r: str(r.get("报告日", "")), reverse=True)
        except Exception:
            return recs

    bal, cf, inc = sheet("balance"), sheet("cashflow"), sheet("income")
    if not bal:
        return out
    latest = bal[0]

    def gv(rec, key):
        return _amt_to_yi((rec or {}).get(key))

    out["report_date"] = str(latest.get("报告日", ""))
    out["assets"] = gv(latest, "资产总计")
    out["liabilities"] = gv(latest, "负债合计")
    out["equity"] = gv(latest, "所有者权益(或股东权益)合计")
    if out["equity"] is None:
        out["equity"] = gv(latest, "所有者权益合计")
    out["cash"] = gv(latest, "货币资金")
    # 商誉科目存在但为空 = 账面无商誉（如实填 0，供排雷比率计算）
    if "商誉" in latest:
        out["gw"] = gv(latest, "商誉")
        if out["gw"] is None:
            out["gw"] = 0.0
    out["inventory"] = gv(latest, "存货")
    out["ar"] = gv(latest, "应收账款")
    ib = [gv(latest, k) for k in ("短期借款", "一年内到期的非流动负债", "长期借款", "应付债券")]
    ib = [x for x in ib if x is not None]
    out["idebt"] = round(sum(ib), 2) if ib else None

    # 同比期（去年同报告日）
    d = str(latest.get("报告日", ""))
    m = re.match(r"(\d{4})(-\d{2}-\d{2})$", d)
    prev_d = f"{int(m.group(1)) - 1}{m.group(2)}" if m else ""
    prev = next((r for r in bal[1:] if str(r.get("报告日", "")) == prev_d), None)
    if prev:
        pinv, par = gv(prev, "存货"), gv(prev, "应收账款")
        try:
            # 基数过小（<1 亿）时同比失真，不填
            if out["inventory"] is not None and pinv and pinv >= 1:
                out["invG"] = round((out["inventory"] / pinv - 1) * 100, 1)
            if out["ar"] is not None and par and par >= 1:
                out["arG"] = round((out["ar"] / par - 1) * 100, 1)
        except Exception:
            pass

    if cf:
        out["ocf"] = gv(cf[0], "经营活动产生的现金流量净额")
        out["capex"] = gv(cf[0], "购建固定资产、无形资产和其他长期资产所支付的现金")
        if out["capex"] is None:
            out["capex"] = gv(cf[0], "购建固定资产、无形资产和其他长期资产支付的现金")
        if out["ocf"] is not None and out.get("capex") is not None:
            out["fcf"] = round(out["ocf"] - out["capex"], 2)

    # 分红率近似：最近年报期「分配股利、利润或偿付利息支付的现金」/ 年报净利润
    # （含利息支付，略偏高；新浪「股息发放率」列数据残缺不可用）
    def _year_sheet(recs):
        return next((r for r in recs if str(r.get("报告日", "")).endswith("12-31")), None)

    cf_y, inc_y = _year_sheet(cf), _year_sheet(inc)
    div_pay = gv(cf_y, "分配股利、利润或偿付利息所支付的现金")
    if div_pay is None:
        div_pay = gv(cf_y, "分配股利、利润或偿付利息支付的现金")
    np_year = gv(inc_y, "净利润")
    if div_pay is not None and np_year and div_pay > 0 and np_year > 0:
        out["payout"] = round(div_pay / np_year * 100, 1)

    def _gm(rec):
        rev, cost = gv(rec, "营业总收入"), gv(rec, "营业成本")
        if rev and cost is not None:
            return (rev - cost) / rev * 100
        return None

    if inc:
        g0 = _gm(inc[0])
        pinc = next((r for r in inc[1:] if str(r.get("报告日", "")) == prev_d), None)
        g1 = _gm(pinc) if pinc else None
        if g0 is not None and g1 is not None:
            out["gmD"] = round(g0 - g1, 2)
    return out

def get_dividend_payout(code: str):
    """最近年报的股息发放率（%，>0），来自新浪年度财务指标；无有效数据返回 None"""
    try:
        import akshare as ak
        df = ak.stock_financial_analysis_indicator(symbol=code, start_year=str(datetime.now().year - 4))
        if df is None or df.empty:
            return None
        for _, row in df[::-1].iterrows():
            v = row.get("股息发放率(%)")
            try:
                fv = float(str(v).replace('%', '').replace(',', ''))
                # 过滤 NaN 与 0（最新期指标常未公布而记 0）
                if fv == fv and fv > 0:
                    return round(fv, 1)
            except (TypeError, ValueError):
                continue
    except Exception:
        pass
    return None

def get_industry_aggregates(industry_name: str, cur_net_profit_yi: float = None) -> dict:
    """巨潮行业聚合快照：按行业名称匹配证监会/中上协分类。

    返回 {industry, report_date, company_count, total_mcap_yi, net_profit_yi, pe_w,
          net_profit_yoy, share_profit}：
    - net_profit_yoy：行业净利润合计同比（本期 vs 去年同期）→ 可作"行业增速"参考口径
    - share_profit：公司净利润 / 行业净利润合计（利润口径市占率近似）
    营收口径的行业规模与 CR3/CR5 无免费公开接口。
    """
    name = (industry_name or '').strip()
    if not name:
        return {}
    now = datetime.now()
    cands = []
    for yy in (now.year, now.year - 1, now.year - 2):
        for md in ('1231', '0930', '0630', '0331'):
            cands.append(f'{yy}{md}')
    try:
        cur_d = next(c for c in cands
                     if datetime.strptime(c, '%Y%m%d') <= now - timedelta(days=60))
    except StopIteration:
        return {}
    prev_d = f"{int(cur_d[:4]) - 1}{cur_d[4:]}"

    def _cninfo_industry_df(date):
        try:
            import akshare as ak
            return ak.stock_industry_pe_ratio_cninfo(symbol="证监会行业分类", date=date)
        except Exception:
            return None

    def _pick(df):
        if df is None or df.empty or '行业名称' not in df.columns:
            return None
        hit = df[df['行业名称'].astype(str) == name]
        if hit.empty and len(name) >= 4:
            hit = df[df['行业名称'].astype(str).str.contains(name[:4], na=False, regex=False)]
        if hit.empty:
            return None

        def _num(v):
            try:
                x = float(v)
                return x if x == x else None
            except (TypeError, ValueError):
                return None

        r = hit.iloc[-1]
        return {
            'industry': str(r.get('行业名称', name)),
            'company_count': _num(r.get('公司数量')),
            'total_mcap_yi': _num(r.get('总市值-静态')),
            'net_profit_yi': _num(r.get('净利润-静态')),
            'pe_w': _num(r.get('静态市盈率-加权平均')),
        }

    cur = _pick(_cninfo_industry_df(cur_d))
    if not cur:
        return {}
    out = dict(cur)
    out['report_date'] = f"{cur_d[:4]}-{cur_d[4:6]}-{cur_d[6:]}"
    prev = _pick(_cninfo_industry_df(prev_d))
    if prev and prev.get('net_profit_yi') and cur.get('net_profit_yi'):
        try:
            out['net_profit_yoy'] = round((cur['net_profit_yi'] / prev['net_profit_yi'] - 1) * 100, 1)
        except Exception:
            pass
    if cur.get('net_profit_yi') and cur_net_profit_yi:
        try:
            out['share_profit'] = round(float(cur_net_profit_yi) / cur['net_profit_yi'] * 100, 1)
        except Exception:
            pass
    return out

# 东财业绩报表缓存（全市场，按报告期；下载约 15-20s，进程内只拉一次）
_YJBB_CACHE = {}

def _yjbb_df(date: str):
    """东财业绩报表（全市场快照，含所处行业/营收/同比）；失败返回 None"""
    if date in _YJBB_CACHE:
        return _YJBB_CACHE[date]
    df = None
    try:
        import akshare as ak
        d = ak.stock_yjbb_em(date=date)
        if d is not None and not d.empty:
            df = d
    except Exception:
        df = None
    _YJBB_CACHE[date] = df
    return df

def _num_or_none(v):
    try:
        x = float(v)
        return x if x == x else None
    except (TypeError, ValueError):
        return None

def _recent_report_date() -> str:
    """最近已完整披露的报告期（季末日后 60 天），格式 YYYYMMDD"""
    now = datetime.now()
    cands = []
    for yy in (now.year, now.year - 1, now.year - 2):
        for md in ('1231', '0930', '0630', '0331'):
            cands.append(f'{yy}{md}')
    try:
        return next(c for c in cands if datetime.strptime(c, '%Y%m%d') <= now - timedelta(days=60))
    except StopIteration:
        return ''

def get_industry_stats_em(code: str) -> dict:
    """东财业绩报表聚合：营收口径的行业规模/增速/市占率/CR3/CR5。

    行业取业绩报表自带的「所处行业」分类（东财行业，如 白酒Ⅱ）。
    行业营收同比用 Σ本期 / Σ(本期/(1+同比%)) 近似重构去年同期（剔除 |同比|≥90% 奇异值）。
    东财接口不可达时返回 {}，由调用方回退巨潮利润口径。
    """
    code = str(code or '').strip()
    if not code:
        return {}
    cur_d = _recent_report_date()
    if not cur_d:
        return {}
    df = _yjbb_df(cur_d)
    if df is None or '所处行业' not in df.columns or '股票代码' not in df.columns:
        return {}
    row = df[df['股票代码'].astype(str) == code]
    if row.empty:
        return {}
    r = row.iloc[-1]
    ind_name = str(r.get('所处行业') or '').strip()
    if not ind_name:
        return {}
    ind = df[df['所处行业'].astype(str) == ind_name].copy()
    rev_col, g_col = '营业总收入-营业总收入', '营业总收入-同比增长'
    ind['_rev'] = ind[rev_col].map(_num_or_none)
    ind = ind[ind['_rev'].notna() & (ind['_rev'] > 0)]
    if ind.empty:
        return {}
    total_rev = float(ind['_rev'].sum())
    out = {
        'industry': ind_name,
        'report_date': f'{cur_d[:4]}-{cur_d[4:6]}-{cur_d[6:]}',
        'company_count': int(len(ind)),
        'industry_rev_yi': round(total_rev / 1e8, 1),
        'share_rev': None, 'company_rev_yi': None, 'rev_yoy': None,
        'cr3': None, 'cr5': None, 'peers_top': [],
    }
    company_rev = _num_or_none(r.get(rev_col))
    if company_rev and company_rev > 0:
        out['share_rev'] = round(company_rev / total_rev * 100, 1)
        out['company_rev_yi'] = round(company_rev / 1e8, 1)
    g = ind[g_col].map(_num_or_none)
    ok = g.notna() & (g.abs() < 90)
    if int(ok.sum()) >= max(3, int(len(ind) * 0.5)):
        prev_total = float((ind.loc[ok, '_rev'] / (1 + g[ok] / 100)).sum())
        if prev_total > 0:
            out['rev_yoy'] = round((total_rev / prev_total - 1) * 100, 1)
    top = ind.nlargest(5, '_rev')
    out['peers_top'] = [str(x) for x in top.get('股票简称', []).tolist()]
    if total_rev > 0 and len(top) >= 3:
        out['cr3'] = round(float(top['_rev'].head(3).sum()) / total_rev * 100, 1)
        if len(top) >= 5:
            out['cr5'] = round(float(top['_rev'].sum()) / total_rev * 100, 1)
    return out


def get_financial_analysis(code: str, market: str) -> dict:
    """财务分析：三表关键数据 + 财务比率 + 杜邦 + 报表绝对值"""
    out = {"ok": False, "ratios": None, "dupont": None, "statements": {}}
    fd = fetch_financial_data(code)
    if not fd or not fd.get('ok'):
        return out
    rr = calculate_ratios(fd)
    if rr and rr.get('ok'):
        out["ok"] = True
        out["ratios"] = rr["ratios"]
        out["dupont"] = rr["ratios"].get("dupont")
    try:
        out["statements"] = _extract_statements(fd)
    except Exception:
        pass
    return out


def get_market_data(code: str) -> dict:
    """市场表现：现价/PE/PB/52周高低/市值（快照 + 历史K线高低）"""
    out = {"price": None, "pe": None, "pb": None, "roe": None, "gross_margin": None,
           "net_margin": None, "debt_ratio": None, "revenue_yi": None, "net_profit_yi": None,
           "market_cap_yi": None, "high_52": None, "low_52": None, "mkt": "A股"}
    sn = _snapshot(code)
    if not sn or not sn.get("price"):
        return out
    out["price"] = sn.get("price")
    out["pe"] = sn.get("pe")
    out["pb"] = sn.get("pb")
    out["roe"] = round(sn["roe"], 2) if sn.get("roe") else None
    out["gross_margin"] = round(sn["gross_margin"], 2) if sn.get("gross_margin") else None
    out["net_margin"] = round(sn["net_margin"], 2) if sn.get("net_margin") else None
    out["debt_ratio"] = round(sn["debt_ratio"], 2) if sn.get("debt_ratio") else None
    out["revenue_yi"] = round(sn["revenue"] / 1e8, 2) if sn.get("revenue") else None
    out["net_profit_yi"] = round(sn["net_profit"] / 1e8, 2) if sn.get("net_profit") else None
    out["mkt"] = "港股" if sn.get("market") == "hk" else "A股"
    if sn.get("shares"):
        try:
            out["market_cap_yi"] = round(float(sn["price"]) * float(sn["shares"]) / 1e8, 1)
        except Exception:
            out["market_cap_yi"] = None
    # 52周高低（近250交易日K线）
    try:
        import akshare as ak
        if out["mkt"] == "A股":
            sina = ('sh' if code.startswith('6') else 'sz') + code
            df = ak.stock_zh_a_daily(symbol=sina, adjust="")
        else:
            df = ak.stock_hk_daily(symbol=code[2:], adjust="")
        if df is not None and not df.empty:
            tail = df.tail(250)
            out["high_52"] = round(float(tail["high"].max()), 2)
            out["low_52"] = round(float(tail["low"].min()), 2)
    except Exception:
        pass
    return out


def get_dcf_valuation(code: str) -> dict:
    """DCF 绝对估值 + 相对估值（可比公司）"""
    out = {"dcf": None, "peer": None}
    dcf = dcf_valuation(code)
    if dcf and not dcf.get("error"):
        out["dcf"] = dcf
    peers = PEERS_BY_CODE.get(code)
    if peers:
        pc = peer_comparison(code, peers=peers)
        if pc and not pc.get("error"):
            out["peer"] = pc
    return out


def get_news(code: str, limit: int = 6) -> list:
    """最新新闻与舆情：东财个股新闻"""
    news = []
    try:
        import akshare as ak
        df = ak.stock_news_em(symbol=code)
        if df is not None and not df.empty:
            cols = [c for c in ['标题', '时间', '文章来源', '新闻内容'] if c in df.columns]
            for _, row in df.head(limit).iterrows():
                item = {}
                item["title"] = str(row.get('标题', '')).strip()
                item["time"] = str(row.get('时间', '')).strip()
                item["source"] = str(row.get('文章来源', '')).strip()
                item["content"] = str(row.get('新闻内容', '')).strip().replace('\n', ' ')
                if item["title"]:
                    news.append(item)
    except Exception:
        pass
    return news


def _auto_peers(code: str) -> list:
    """`PEERS_BY_CODE` 未覆盖时，按公司名识别行业后由 `peer_pool` 补池。

    池外行业原来会**静默退化为空**（可比公司为空，可比估值失去参照）；
    现在改为按行业池补全，且只保留能取到行情的代码（A 股 / 港股），不臆造同业。
    """
    try:
        from peer_pool import resolve_peers, is_fetchable_code
        name = code_to_name(code) if callable(code_to_name) else ''
        dec = resolve_peers(code=code, name=name or '', industry_hint=name or '')
        return [p['code'] for p in (dec.get('peers') or [])
                if isinstance(p, dict) and is_fetchable_code(p.get('code'))]
    except Exception as e:
        logger.debug('同业池自动补全失败: %s', e)
        return []


def get_competitors(code: str) -> list:
    """竞争对手对比：同行业可比公司列表（并行取快照，见 `market_snapshot`）。

    原实现串行 `for p in peers: _snapshot(p)`，实测 **30.06s**（茅台，2026-10-02），
    是全管线最慢一步；改为并行后与 `peer_comparison` 同策略。
    """
    from market_snapshot import competitor_rows
    peers = PEERS_BY_CODE.get(code) or _auto_peers(code)
    return competitor_rows(code, peers, code_to_name)


# ==================== SWOT 与风险（规则判定，基于真实财务/市场数据） ====================

def generate_swot(market: dict, fin: dict, profile: dict = None, dcf: dict = None) -> dict:
    """基于真实数据生成 SWOT，规则化判定（避免臆断）。

    实现已收敛到 `company_deep_narrative.generate_swot`（单一实现）：
    本函数原先接受 `profile` 与 `dcf` 两个入参却**从未使用**，
    "行业位置/行业份额"与"DCF 内在价值 vs 现价"两个维度实为空缺。
    `profile`/`dcf` 现改为可选，缺省时相应维度自然略过。
    """
    from company_deep_narrative import generate_swot as _impl
    return _impl(market, fin, profile, dcf)


def generate_risks(market: dict, fin: dict) -> list:
    """风险提示：基于真实数据归纳"""
    risks = []
    dr = market.get("debt_ratio")
    gm = market.get("gross_margin")
    rev_g = fin.get("rev_growth")
    pro_g = fin.get("profit_growth")
    pe = market.get("pe")

    if dr is not None and dr >= 60:
        risks.append({"level": "高", "name": "财务杠杆风险",
                      "desc": f"资产负债率达 {dr}%，偿债与再融资压力较大。若利率上行或融资环境收紧，财务费用将侵蚀利润。"})
    elif dr is not None and dr >= 40:
        risks.append({"level": "中", "name": "财务杠杆风险", "desc": f"资产负债率 {dr}%，财务结构尚可但需关注负债变化趋势。"})
    if pro_g is not None and pro_g < 0:
        risks.append({"level": "高", "name": "盈利下滑风险", "desc": f"净利润同比下滑 {pro_g}%，成长动能减弱。"})
    if rev_g is not None and rev_g < 0:
        risks.append({"level": "中", "name": "需求放缓风险", "desc": f"营收同比下滑 {rev_g}%，行业景气度承压。"})
    if pe is not None and pe and pe >= 40:
        risks.append({"level": "高", "name": "估值风险", "desc": f"PE {pe} 估值偏高，隐含较高增长预期，业绩不及预期时回撤显著。"})
    if gm is not None and gm < 15:
        risks.append({"level": "中", "name": "竞争议价风险", "desc": f"毛利率 {gm}% 偏低，产品或面临激烈竞争、成本上升压力。"})
    if not risks:
        risks.append({"level": "低", "name": "综合风险", "desc": "当前数据未显示显著风险，但仍需关注行业政策、宏观经济与竞争格局变化。"})
    return risks


# ==================== 综合分析入口 ====================

def analyze_company(company_name: str):
    """公司全方位分析主入口：聚合各维度数据 → 返回结构化 report dict"""
    code, market = resolve_ticker(company_name)
    if not code:
        return {"error": f"无法识别公司：{company_name}。请提供公司名（如 贵州茅台/腾讯）或股票代码（如 600519/HK0700）。"}

    profile = get_company_profile(code, market)
    name = code_to_name(code)
    if name == code:
        name = profile["name"] or code
    profile["name"] = name

    # 财务分析
    fa = get_financial_analysis(code, market)
    ratios = fa.get("ratios") if fa.get("ok") else {}
    fin = {
        "rev_growth": (ratios.get("growth") or {}).get("revenue_growth"),
        "profit_growth": (ratios.get("growth") or {}).get("profit_growth"),
        "rev_yi": (ratios.get("profitability") or {}).get("revenue"),
        "net_profit_yi": (ratios.get("profitability") or {}).get("net_profit"),
        "eps": (ratios.get("profitability") or {}).get("eps"),
        "roe": (ratios.get("profitability") or {}).get("roe"),
        "dupont": ratios.get("dupont"),
    }

    # 市场表现
    market_data = get_market_data(code)

    # 估值
    val = get_dcf_valuation(code)
    dcf, peer = val.get("dcf"), val.get("peer")

    # 新闻
    news = get_news(code)

    # 竞争对手
    competitors = get_competitors(code)

    # SWOT + 风险（规则）
    swot = generate_swot(market_data, fin, profile, val)
    risks = generate_risks(market_data, fin)

    report = {
        "ok": True,
        "name": name,
        "code": code,
        "analyzed_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "profile": profile,
        "financial": fin,
        "market": market_data,
        "valuation": {"dcf": dcf, "peer": peer},
        "news": news or [],
        "competitors": competitors,
        "swot": swot,
        "risks": risks,
    }
    return report


# ==================== Word 报告生成（python-docx） ====================

def _add_kv_table(doc, rows):
    """key-value 表格"""
    table = doc.add_table(rows=1, cols=2)
    try:
        table.style = 'Light Grid Accent 1'
    except Exception:
        pass
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = '指标', '数值'
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v if v is not None else '-')
    doc.add_paragraph("")


def generate_docx(report: dict) -> bytes:
    """生成 Word 报告，返回 docx bytes"""
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    name = report.get("name", "")
    doc.add_heading(f"{name} 全方位分析报告", 0)
    doc.add_paragraph(f"分析时间：{report.get('analyzed_at', '')}    股票代码：{report.get('code', '')}")
    doc.add_paragraph("数据来源：akshare 公开金融数据接口（A股/港股）。仅供研究参考，不构成投资建议。")

    # 一、公司概况
    doc.add_heading("一、公司概况", 1)
    p = report.get("profile", {})
    _add_kv_table(doc, [
        ("公司名称", p.get("name")), ("所属行业", p.get("industry")),
        ("总市值", f"{p.get('market_cap_yi')} 亿元" if p.get('market_cap_yi') else None),
        ("总股本", f"{p.get('total_share_yi')} 亿股" if p.get('total_share_yi') else None),
        ("上市时间", p.get("list_time")),
    ])

    # 二、财务分析
    doc.add_heading("二、财务分析", 1)
    f = report.get("financial", {})
    _add_kv_table(doc, [
        ("营业收入(亿)", f.get("rev_yi")), ("净利润(亿)", f.get("net_profit_yi")),
        ("营收同比", f"{f.get('rev_growth')}%" if f.get('rev_growth') is not None else None),
        ("净利同比", f"{f.get('profit_growth')}%" if f.get('profit_growth') is not None else None),
        ("EPS(元)", f.get("eps")), ("ROE(%)", f.get("roe")),
    ])
    dup = f.get("dupont")
    if dup:
        doc.add_heading("杜邦分析", 2)
        _add_kv_table(doc, [
            ("ROE(%)", dup.get("roe")), ("净利率(%)", dup.get("net_margin")),
            ("总资产周转率", dup.get("asset_turnover")), ("权益乘数", dup.get("equity_multiplier")),
        ])
        doc.add_paragraph("杜邦拆解：ROE = 净利率 × 总资产周转率 × 权益乘数，可定位盈利能力来源。", style="Intense Quote")


    # 三、市场表现
    doc.add_heading("三、市场表现", 1)
    m = report.get("market", {})
    _add_kv_table(doc, [
        ("当前股价(元)", m.get("price")), ("市场", m.get("mkt")),
        ("总市值(亿)", m.get("market_cap_yi")), ("PE", m.get("pe")),
        ("PB", m.get("pb")), ("ROE(%)", m.get("roe")),
        ("毛利率(%)", m.get("gross_margin")), ("净利率(%)", m.get("net_margin")),
        ("资产负债率(%)", m.get("debt_ratio")), ("52周最高", m.get("high_52")),
        ("52周最低", m.get("low_52")),
    ])

    # 四、估值
    doc.add_heading("四、估值分析", 1)
    dcf = report.get("valuation", {}).get("dcf")
    if dcf:
        doc.add_heading("DCF 绝对估值", 2)
        _add_kv_table(doc, [
            ("现价(元)", dcf.get("price")), ("每股内在价值(元)", dcf.get("per_share")),
            ("偏差(%)", dcf.get("spread_pct")), ("结论", dcf.get("verdict_cn")),
        ])
    peer = report.get("valuation", {}).get("peer")
    if peer and peer.get("verdict"):
        doc.add_heading("相对估值（可比公司）", 2)
        doc.add_paragraph(peer["verdict"]["label"])
        rows = peer.get("peers", [])
        if rows:
            table = doc.add_table(rows=1, cols=5)
            try:
                table.style = 'Light Grid Accent 1'
            except Exception:
                pass
            hdr = table.rows[0].cells
            for i, h in enumerate(["代码", "名称", "PE", "PB", "ROE(%)"]):
                hdr[i].text = h
            for r in rows:
                cells = table.add_row().cells
                cells[0].text = str(r.get("code", "-"))
                cells[1].text = str(r.get("name", "-"))
                cells[2].text = str(r.get("pe", "-"))
                cells[3].text = str(r.get("pb", "-"))
                cells[4].text = str(r.get("roe", "-"))
            doc.add_paragraph("")

    # 五、最新新闻与舆情
    doc.add_heading("五、最新新闻与舆情", 1)
    news = report.get("news", [])
    if news:
        for i, n in enumerate(news[:6], 1):
            doc.add_paragraph(f"{i}. {n.get('time','')} [{n.get('source','')}] {n.get('title','')}", style="List Number")
            if n.get("content"):
                doc.add_paragraph(f"   {n['content'][:120]}", style="Intense Quote")
    else:
        doc.add_paragraph("未获取到相关新闻（数据源暂时不可用或无覆盖）。")

    # 六、竞争对手对比
    doc.add_heading("六、竞争对手对比", 1)
    comps = report.get("competitors", [])
    if comps:
        for c in comps:
            doc.add_paragraph(f"{c.get('name')}（{c.get('code')}）")
    else:
        doc.add_paragraph("暂无已配置的同行业竞争对手映射。")

    # 七、SWOT 分析
    doc.add_heading("七、SWOT 分析", 1)
    sw = report.get("swot", {})
    for label, key in [("优势 (Strengths)", "strengths"), ("劣势 (Weaknesses)", "weaknesses"),
                        ("机会 (Opportunities)", "opportunities"), ("威胁 (Threats)", "threats")]:
        doc.add_heading(label, 2)
        for item in sw.get(key, []):
            doc.add_paragraph(f"• {item}")

    # 八、风险提示
    doc.add_heading("八、风险提示", 1)
    table = doc.add_table(rows=1, cols=3)
    try:
        table.style = 'Light Grid Accent 1'
    except Exception:
        pass
    hdr = table.rows[0].cells
    for i, h in enumerate(["风险等级", "风险项", "说明"]):
        hdr[i].text = h
    for r in report.get("risks", []):
        doc_row = table.add_row().cells
        doc_row[0].text = str(r.get("level", "-"))
        doc_row[1].text = str(r.get("name", "-"))
        doc_row[2].text = str(r.get("desc", "-"))
    doc.add_paragraph("")

    # 免责声明
    para = doc.add_paragraph()
    para.add_run("\n免责声明：本报告基于 akshare 公开数据自动生成，数据可能存在延迟或误差，仅供研究参考，不构成任何投资建议。投资有风险，入市需谨慎。")
    para.runs[0].font.size = Pt(9)
    para.runs[0].font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_docx_report_file(report: dict) -> str:
    """生成 docx 并保存到 output 目录，返回文件路径"""
    data = generate_docx(report)
    fname = f"{report.get('name')}_{datetime.now().strftime('%Y%m%d_%H%M%S')}_analysis.docx"
    path = OUT_DIR / fname
    path.write_bytes(data)
    return str(path)