#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
company_deep_ppt.py — 企业深析 PPT 导出（版式编排，从 company_deep_report 拆出）

绘制组件在 company_deep_ppt_kit，主题在 company_deep_themes，指标在 company_deep_metrics。
"""
import io
from datetime import datetime

from company_deep_metrics import _s, _n, _fmt, derive_metrics, build_verdict_lines
from company_deep_themes import PPT_TEMPLATES, DEFAULT_PPT_TEMPLATE, list_ppt_templates, resolve_theme  # noqa: F401
from company_deep_ppt_kit import (  # noqa: F401
    _add_text, _add_shape, _add_rect, _set_alpha, _gradient_rect,
    _mix, _soften, _faint, _is_dark, _card_color, _hi,
    _slide_title, _content_chrome, _stat_card, _pill, _num_badge,
    _bar_row, _note_row, _text_card, _bullet_block, _paint_slide_bg,
)

def generate_pptx(state: dict, template: str | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    th = resolve_theme(template)
    ACCENT, ACCENT2 = th["accent"], th["accent2"]
    OK, WARN, DANGER = th["ok"], th["warn"], th["danger"]
    INK, MUTED = th["ink"], th["muted"]
    ON_DARK, ON_DARK_SUB = th["on_dark"], th["on_dark_sub"]
    CARD = _card_color(th)
    BG = th.get("slide_bg") or (0xFF, 0xFF, 0xFF)
    DEEP = _mix(ACCENT, (0, 0, 0), 0.45)

    met = derive_metrics(state)
    verdicts = build_verdict_lines(state, met)
    c = state.get("company") or {}
    name = _s(c.get("name"), "未命名企业")
    now = datetime.now().strftime("%Y-%m-%d")
    tag = f"{name}  ·  {_s(c.get('code'))}"

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    def set_notes(s, text):
        try:
            s.notes_slide.notes_text_frame.text = text
        except Exception:
            pass

    def pill_row(s, x, y, pairs, th_):
        """[(文字, 颜色)] 一排标签"""
        cx = x
        for txt, col in pairs:
            cx = _pill(s, cx, y, txt, th_, accent=col)

    # ── 1 封面 ──
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if th.get("cover_mode") == "minimal":
        _add_shape(s, "RECTANGLE", 0, 0, 13.333, 7.5, (0xFF, 0xFF, 0xFF))
        _add_shape(s, "RECTANGLE", 0, 0, 0.14, 7.5, ACCENT)
        # 右上装饰圆
        d = _add_shape(s, "OVAL", 10.4, -1.6, 4.4, 4.4, _mix(ACCENT2, (0xFF, 0xFF, 0xFF), 0.86))
        d2 = _add_shape(s, "OVAL", 11.6, 2.3, 2.4, 2.4, _soften(th, 0.80))
        _add_text(s, 1.0, 1.55, 10.5, 0.4, "ENTERPRISE DEEP DIVE  ·  企业深度分析", size=13,
                  bold=True, color=ACCENT2)
        _add_text(s, 1.0, 2.15, 11.0, 1.15, name, size=54, bold=True, color=INK)
        _add_text(s, 1.0, 3.45, 11.0, 0.45,
                  f"{_s(c.get('code'))}    ·    {_s(c.get('industry'))}    ·    {_s(c.get('market'))}",
                  size=17, color=MUTED)
        _add_text(s, 1.0, 4.15, 11.0, 0.6, _s(c.get("oneliner"), "漏斗式研判 · 五问收口"),
                  size=15, color=INK)
        _add_shape(s, "RECTANGLE", 1.0, 5.35, 10.9, 0.014, _faint(th, 0.85))
        _add_text(s, 1.0, 5.55, 11.0, 0.4,
                  f"分析日期 {_s(c.get('date'), now)}    ·    分析人 {_s(c.get('analyst'))}    ·    data.tool 企业深析",
                  size=12, color=MUTED)
    elif th.get("cover_mode") == "band":
        _add_shape(s, "RECTANGLE", 0, 0, 13.333, 7.5, (0xFF, 0xFF, 0xFF) if th["panel"][0] > 200 else th["panel"])
        _gradient_rect(s, 0, 0, 13.333, 1.12, ACCENT, DEEP, angle=0)
        _add_text(s, 0.85, 0.36, 11.5, 0.45, "企业深度分析报告", size=15, bold=True, color=ON_DARK)
        _add_shape(s, "OVAL", 11.15, 2.5, 3.6, 3.6, _soften(th, 0.88))
        _add_shape(s, "PARALLELOGRAM", -0.4, 6.1, 6.5, 0.55, _soften(th, 0.75), adj=0.6)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.85, 2.05, 0.62, 0.06, ACCENT2, adj=0.5)
        _add_text(s, 0.85, 2.35, 11.0, 1.1, name, size=52, bold=True, color=INK)
        _add_text(s, 0.85, 3.6, 11.0, 0.45,
                  f"{_s(c.get('code'))}    ·    {_s(c.get('industry'))}    ·    {_s(c.get('market'))}",
                  size=17, color=MUTED)
        _add_text(s, 0.85, 4.3, 11.0, 0.6, _s(c.get("oneliner"), "漏斗式研判 · 五问收口"), size=15, color=INK)
        _add_text(s, 0.85, 5.55, 11.0, 0.4,
                  f"分析日期 {_s(c.get('date'), now)}    ·    分析人 {_s(c.get('analyst'))}    ·    模板 {th['name']}",
                  size=12, color=MUTED)
    else:  # solid 渐变封面
        _gradient_rect(s, 0, 0, 13.333, 7.5, ACCENT, DEEP, angle=115)
        d = _add_shape(s, "OVAL", 9.6, -2.0, 5.6, 5.6, ACCENT2)
        _set_alpha(d, 10)
        d2 = _add_shape(s, "OVAL", 11.3, 3.6, 3.4, 3.4, (0xFF, 0xFF, 0xFF))
        _set_alpha(d2, 6)
        p1 = _add_shape(s, "PARALLELOGRAM", -0.5, 5.9, 7.4, 0.62, ACCENT2, adj=0.6)
        _set_alpha(p1, 26)
        p2 = _add_shape(s, "PARALLELOGRAM", 1.6, 6.65, 7.4, 0.62, (0xFF, 0xFF, 0xFF), adj=0.6)
        _set_alpha(p2, 10)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.85, 1.72, 2.35, 0.44, ACCENT2, adj=0.5)
        _add_text(s, 0.85, 1.79, 2.35, 0.32, "企业深度分析报告", size=12.5, bold=True, color=DEEP,
                  align=PP_ALIGN.CENTER)
        _add_text(s, 0.85, 2.5, 11.6, 1.15, name, size=54, bold=True, color=ON_DARK)
        _add_text(s, 0.85, 3.82, 11.6, 0.45,
                  f"{_s(c.get('code'))}    ·    {_s(c.get('industry'))}    ·    {_s(c.get('market'))}",
                  size=17, color=ON_DARK_SUB)
        _add_text(s, 0.85, 4.55, 11.6, 0.6, _s(c.get("oneliner"), "漏斗式研判 · 五问收口"),
                  size=15, color=ON_DARK)
        _add_shape(s, "RECTANGLE", 0.85, 5.62, 11.0, 0.012, ON_DARK_SUB)
        _add_text(s, 0.85, 5.82, 11.6, 0.4,
                  f"分析日期 {_s(c.get('date'), now)}    ·    分析人 {_s(c.get('analyst'))}    ·    data.tool 企业深析",
                  size=12, color=ON_DARK_SUB)
    set_notes(s, "开场：说明本报告按漏斗框架完成，重点看生意、公司、账、风险、价格五问。")

    # ── 2 一页纸判断卡 ──
    s = _content_chrome(prs, th, 1, "一页纸判断卡", "五问收口：生意 / 公司 / 账 / 风险 / 价格", tag, 2)
    y = 1.42
    for i, (q, a, _ev) in enumerate(verdicts[:5]):
        _add_rect(s, 0.55, y, 12.23, 0.92, CARD)
        _num_badge(s, 0.78, y + 0.25, i + 1, th)
        _add_text(s, 1.42, y + 0.13, 2.5, 0.3, q, size=12.5, bold=True, color=_hi(th, ACCENT))
        _add_text(s, 1.42, y + 0.46, 11.1, 0.4, a if len(a) < 92 else a[:90] + "…", size=11.5, color=INK)
        y += 1.06
    set_notes(s, "逐条读五问结论；细节在后续分页展开。")

    # ── 3 赛道 ──
    s = _content_chrome(prs, th, 2, "赛道决定天花板", "先看赛道，再选马", tag, 3)
    b = state.get("business") or {}
    _stat_card(s, 0.55, 1.35, 2.92, 1.32, _s(b.get("size")), "行业规模（亿元）", ACCENT, th)
    _stat_card(s, 3.65, 1.35, 2.92, 1.32, _s(b.get("growth")) + "%", "行业增速（%）",
               OK if (_n(b.get("growth")) or 0) >= 0 else DANGER, th)
    _stat_card(s, 6.75, 1.35, 2.92, 1.32, _s(b.get("share")) + "%", "公司市占率（%）", ACCENT2, th)
    _stat_card(s, 9.85, 1.35, 2.92, 1.32, f"{_s(b.get('score'), '—')}/5", "行业评分", WARN, th)
    pill_row(s, 0.55, 2.98, [
        (f"生命周期 {_s(b.get('lifecycle'))}", ACCENT2),
        (f"政策 {_s(b.get('policy'))}", ACCENT),
        (f"CR3 {_s(b.get('cr3'))}%", WARN),
        (f"CR5 {_s(b.get('cr5'))}%", WARN),
    ], th)
    _text_card(s, 0.55, 3.62, 6.02, 1.52, "靠什么赚钱", _s(b.get("what"), "—"), th, ACCENT)
    _text_card(s, 6.75, 3.62, 6.02, 1.52, "行业天花板", _s(b.get("ceiling"), "—"), th, ACCENT2)
    _text_card(s, 0.55, 5.3, 6.02, 1.52, "关键成功要素", _s(b.get("kfs"), "—"), th, OK)
    _text_card(s, 6.75, 5.3, 6.02, 1.52, "行业位置", _s(b.get("position"), "—"), th, WARN)
    set_notes(s, "强调行业增速与集中度；政策方向直接影响中期空间。")

    # ── 4 商业闭环 ──
    s = _content_chrome(prs, th, 3, "商业闭环：钱从哪来、货往哪流", "找出链条上最脆弱的一环", tag, 4)
    model = state.get("model") or {}
    nodes = [("① 客户", "customer"), ("② 价值", "value"), ("③ 渠道", "channel"),
             ("④ 收入", "revenue"), ("⑤ 成本", "cost"), ("⑥ 利润回款", "cash")]
    for i, (label, key) in enumerate(nodes):
        col, row = i % 3, i // 3
        x = 0.55 + col * 4.22
        y = 1.38 + row * 1.62
        _text_card(s, x, y, 3.9, 1.44, label, _s(model.get(key), "待补充"), th,
                   ACCENT2 if i % 2 == 0 else ACCENT, body_size=11)
        if col < 2:
            ar = _add_shape(s, "RIGHT_ARROW", x + 3.92, y + 0.58, 0.3, 0.26, _soften(th, 0.55))
    _add_shape(s, "RECTANGLE", 0.55, 4.78, 12.23, 0.012, _faint(th))
    _note_row(s, 0.55, 5.0, 12.23, "最脆弱一环", _s(model.get("weakest"), "—"), th, DANGER, h=0.6)
    _text_card(s, 0.55, 5.82, 12.23, 1.05, "利润公式与回款",
               f"利润公式：{_s(model.get('formula'), '—')}    |    回款模式：{_s(model.get('payment'), '—')}"
               f"    |    定价权：{_s(model.get('pricing'), '—')}/5    |    客户集中度：{_s(model.get('cust1'))}% / {_s(model.get('cust5'))}%",
               th, ACCENT)
    set_notes(s, "用大白话讲闭环；指出最脆弱一环及其被攻击方式。")

    # ── 5 护城河 ──
    s = _content_chrome(prs, th, 4, f"护城河 {_fmt(met['moat_avg'])} / 10", "高毛利必引来竞争——凭什么守得住？", tag, 5)
    m = state.get("moat") or {}
    dims = [("技术/专利", "tech"), ("品牌心智", "brand"), ("渠道网络", "channel"),
            ("成本/规模", "cost"), ("网络效应", "network"), ("牌照/资质", "license")]
    y = 1.5
    for label, key in dims:
        val = _n(m.get(key)) or 5.0
        _bar_row(s, 0.55, y, 7.0, label, val / 10.0, f"{val:.0f}", th, ACCENT2)
        y += 0.52
    # 右侧总分卡
    _add_rect(s, 8.0, 1.5, 4.78, 2.5, CARD)
    _add_text(s, 8.0, 1.75, 4.78, 0.9, f"{_fmt(met['moat_avg'])}", size=44, bold=True, color=ACCENT,
              align=PP_ALIGN.CENTER)
    _add_text(s, 8.0, 2.75, 4.78, 0.4, "六维护城河总分 / 10", size=12, color=MUTED, align=PP_ALIGN.CENTER)
    _text_card(s, 0.55, 4.75, 6.02, 1.0, "客户为何非选它", _s(m.get("whyChoose"), "—"), th, ACCENT)
    _text_card(s, 6.75, 4.75, 6.02, 1.0, "对手为何难进来", _s(m.get("whyBlocked"), "—"), th, ACCENT2)
    _text_card(s, 0.55, 5.9, 6.02, 0.95, "市占率变化 / 毛利率 vs 同行",
               f"{_s(m.get('shareTrend'), '—')}    |    {_s(m.get('marginPeers'), '—')}", th, OK, body_size=10.5)
    _text_card(s, 6.75, 5.9, 6.02, 0.95, "壁垒可持续性",
               f"评分 {_s(m.get('dur'), '—')}/5", th, WARN)

    # ── 6 财务验证 ──
    s = _content_chrome(prs, th, 5, "财务反向验证生意判断", "现金流是整条链的照妖镜", tag, 6)
    _stat_card(s, 0.55, 1.35, 2.92, 1.3, _fmt(met["gm"], 1, "%"), "毛利率", ACCENT, th)
    _stat_card(s, 3.65, 1.35, 2.92, 1.3, _fmt(met["nm"], 1, "%"), "净利率", ACCENT2, th)
    _stat_card(s, 6.75, 1.35, 2.92, 1.3, _s((state.get("finance") or {}).get("roe")), "ROE（%）", OK, th)
    _stat_card(s, 9.85, 1.35, 2.92, 1.3, _fmt(met["ocf_np"], 2), "OCF / 净利润",
               OK if (met["ocf_np"] or 0) >= 1 else DANGER, th)
    # 杜邦公式条
    _add_rect(s, 0.55, 2.95, 12.23, 1.0, CARD)
    _add_text(s, 0.75, 3.06, 11.8, 0.4,
              f"ROE ≈ {_fmt(met['du_roe'])}%   =   净利率 {_fmt(met['nm'])}%   ×   周转率 {_fmt(met['tat'], 2)}   ×   权益乘数 {_fmt(met['em'], 2)}",
              size=17, bold=True, color=_hi(th, ACCENT))
    _add_text(s, 0.75, 3.52, 11.8, 0.34,
              "杜邦拆解定位 ROE 驱动源：高净利率（暴利/品牌）· 高周转（薄利多销）· 高杠杆（负债驱动）",
              size=10.5, color=MUTED)
    _add_text(s, 0.55, 4.2, 12.23, 0.35, "现金流排雷扫描", size=13, bold=True, color=INK)
    alerts = met["alerts"] or [("未触发主要排雷信号", "仍需结合多年趋势与应收/存货明细交叉验证")]
    ay = 4.62
    for t, d in alerts[:5]:
        col = DANGER if ("背离" in t or "双高" in t or "过高" in t) else WARN
        _note_row(s, 0.55, ay, 12.23, t, d, th, col, h=0.5, title_w=2.4)
        ay += 0.58
    set_notes(s, "重点解释 ROE 驱动；若 OCF/NP<1 必须提示应收存货。")

    # ── 7 治理与风险 ──
    s = _content_chrome(prs, th, 6, "谁在管？会不会突然暴雷？", "治理决定前面成果能否守住", tag, 7)
    g = state.get("gov") or {}
    _stat_card(s, 0.55, 1.35, 3.4, 1.22, _s(g.get("controller"), "—"), "实际控制人", ACCENT, th)
    _stat_card(s, 4.15, 1.35, 2.6, 1.22, _s(g.get("pledge"), "—"), "质押率（%）", WARN, th)
    _stat_card(s, 6.95, 1.35, 2.6, 1.22, f"{_s(g.get('score'), '—')}/5", "治理评分", OK, th)
    _stat_card(s, 9.75, 1.35, 3.02, 1.22, str(len(state.get("risks") or [])), "已登记风险条数", DANGER, th)
    risks = state.get("risks") or []
    _add_text(s, 0.55, 2.85, 12.23, 0.35, "主要风险", size=13, bold=True, color=INK)
    if risks:
        ry = 3.28
        for r in risks[:6]:
            lv = _s(r.get("level"), "中")
            col = DANGER if lv == "高" else (WARN if lv == "中" else OK)
            _note_row(s, 0.55, ry, 12.23, f"{lv}风险 · {_s(r.get('type'))}", _s(r.get("text"), ""), th, col,
                      h=0.52, title_w=2.0)
            ry += 0.58
    else:
        _note_row(s, 0.55, 3.28, 12.23, "待登记", "风险因素尚未系统登记，建议从招股书「风险因素」逐条提炼", th, WARN, h=0.52)
    _text_card(s, 0.55, 6.35 if risks else 5.9, 12.23, 0.62, "分红与激励",
               f"分红：{_s(g.get('dividend'), '—')}    |    激励：{_s(g.get('incentive'), '—')}",
               th, ACCENT2, body_size=10.5)

    # ── 8 估值 ──
    s = _content_chrome(prs, th, 7, "好公司也要好价格", "估值分位 + 可比 + DCF 锚", tag, 8)
    v = state.get("val") or {}
    _stat_card(s, 0.55, 1.35, 2.92, 1.3, _s(v.get("pe")), "PE（TTM）", ACCENT, th)
    _stat_card(s, 3.65, 1.35, 2.92, 1.3, _s(v.get("pePct")) + "%", "历史 PE 分位", WARN, th)
    _stat_card(s, 6.75, 1.35, 2.92, 1.3, _s(v.get("pb")), "PB", ACCENT2, th)
    _stat_card(s, 9.85, 1.35, 2.92, 1.3, _s(v.get("verdict"), "—"), "价格结论",
               OK if _s(v.get("verdict")) in ("低估", "合理") else DANGER, th)
    # 估值分位条（绿→黄→红 三段 + 标记）
    pct = _n(v.get("pePct"))
    if pct is not None:
        px = 0.55 + 12.23 * max(0.0, min(100.0, pct)) / 100.0
        _add_shape(s, "ROUNDED_RECTANGLE", 0.55, 3.05, 12.23 * 0.5, 0.3, _mix(OK, BG, 0.55), adj=0.5)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.55 + 12.23 * 0.5, 3.05, 12.23 * 0.3, 0.3, _mix(WARN, BG, 0.55), adj=0.5)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.55 + 12.23 * 0.8, 3.05, 12.23 * 0.2, 0.3, _mix(DANGER, BG, 0.55), adj=0.5)
        _add_shape(s, "RECTANGLE", px - 0.02, 2.92, 0.045, 0.56, INK)
        _add_text(s, 0.55, 3.44, 4.0, 0.3, "低估 0%", size=9.5, color=MUTED)
        _add_text(s, 10.8, 3.44, 1.98, 0.3, "高估 100%", size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)
    _text_card(s, 0.55, 3.95, 6.02, 1.35, "相对估值",
               f"同行 PE 中位：{_s(v.get('peerPe'))}    |    PS：{_s(v.get('ps'))}    |    PB 分位：{_s(v.get('pbPct'))}%",
               th, ACCENT2)
    _text_card(s, 6.75, 3.95, 6.02, 1.35, "DCF 参数与结果",
               f"FCF {_s(v.get('dcfFcf'))} 亿 · 永续 g {_s(v.get('dcfG'))}% · WACC {_s(v.get('dcfWacc'))}% · 股本 {_s(v.get('dcfShares'))} 亿股 · 净现金 {_s(v.get('dcfNetcash'))} 亿",
               th, ACCENT, body_size=10.5)
    _text_card(s, 0.55, 5.45, 12.23, 1.4, "估值逻辑（证据链）", _s(v.get("reason"), "—"), th, WARN)
    set_notes(s, "结论必须明确低估/合理/高估，并给出证据链。")

    # ── 9 竞争推演 ──
    pb = state.get("pb") or {}
    if any(pb.get(k) for k in ("strength", "weak", "rival", "conclusion")) or pb.get("steps"):
        s = _content_chrome(prs, th, 8, "站在对手立场推演下一步", "对手全力进攻时守不守得住", tag, 9)
        _text_card(s, 0.55, 1.4, 6.02, 2.6, "我方视角",
                   f"核心优势：{_s(pb.get('strength'), '—')}\n软肋：{_s(pb.get('weak'), '—')}\n下一步：{_s(pb.get('next'), '—')}",
                   th, ACCENT, body_size=11.5)
        _text_card(s, 6.75, 1.4, 6.02, 2.6, "对手视角",
                   f"最大对手：{_s(pb.get('rival'), '—')}\n可能打法：{_s(pb.get('move'), '—')}\n冲击评估：{_s(pb.get('impact'), '—')}",
                   th, DANGER, body_size=11.5)
        steps = pb.get("steps") or []
        if steps:
            _add_text(s, 0.55, 4.25, 12.23, 0.35, "推演步骤", size=13, bold=True, color=INK)
            sy = 4.72
            _add_shape(s, "RECTANGLE", 0.72, sy + 0.1, 0.022, 0.52 * min(len(steps), 4) - 0.2, _faint(th, 0.75))
            for i, st_ in enumerate(steps[:4]):
                _add_shape(s, "OVAL", 0.63, sy + 0.02, 0.2, 0.2, ACCENT2)
                _add_text(s, 1.0, sy - 0.04, 11.7, 0.4, f"Step {i + 1}    {st_}", size=11.5, color=INK)
                sy += 0.52
        _note_row(s, 0.55, 6.35, 12.23, "结论",
                  f"{_s(pb.get('conclusion'), '—')}    |    投资含义：{_s(pb.get('invest'), '—')}",
                  th, ACCENT, h=0.58, title_w=1.0)

    # ── 10 收口 ──
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if th.get("cover_mode") == "minimal":
        _add_shape(s, "RECTANGLE", 0, 0, 13.333, 7.5, (0xFF, 0xFF, 0xFF))
        _add_shape(s, "RECTANGLE", 0, 0, 0.14, 7.5, ACCENT)
        _add_text(s, 0.95, 0.9, 11.2, 0.7, "最终结论 · 五问收口", size=30, bold=True, color=INK)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.98, 1.75, 0.62, 0.055, ACCENT2, adj=0.5)
        y = 2.15
        for i, (q, a, _ev) in enumerate(verdicts):
            _num_badge(s, 0.95, y + 0.06, i + 1, th, d=0.36)
            _add_text(s, 1.5, y, 11.0, 0.75, f"{q}    {a if len(a) < 100 else a[:98] + '…'}",
                      size=12.5, color=INK, line_spacing=1.05)
            y += 0.86
        _add_text(s, 0.95, 6.65, 11.5, 0.35, f"{name}  ·  {now}  ·  {th['name']}  ·  data.tool 企业深析",
                  size=11, color=MUTED)
    else:
        _gradient_rect(s, 0, 0, 13.333, 7.5, ACCENT, DEEP, angle=115)
        d = _add_shape(s, "OVAL", 10.2, 4.6, 4.6, 4.6, ACCENT2)
        _set_alpha(d, 10)
        _add_text(s, 0.85, 0.85, 11.5, 0.7, "最终结论 · 五问收口", size=30, bold=True, color=ON_DARK)
        _add_shape(s, "ROUNDED_RECTANGLE", 0.88, 1.72, 0.62, 0.055, ACCENT2, adj=0.5)
        y = 2.1
        for i, (q, a, _ev) in enumerate(verdicts):
            _num_badge(s, 0.85, y + 0.04, i + 1, th, d=0.36, accent=ACCENT2)
            _add_text(s, 1.4, y, 11.4, 0.75, f"{q}    {a if len(a) < 100 else a[:98] + '…'}",
                      size=12.5, color=ON_DARK, line_spacing=1.05)
            y += 0.86
        _add_text(s, 0.85, 6.45, 11.5, 0.35, f"{name}  ·  {now}  ·  模板 {th['name']}  ·  data.tool 企业深析",
                  size=11, color=ON_DARK_SUB)
        _add_text(s, 0.85, 6.85, 11.5, 0.3,
                  "免责声明：本报告由 data.tool 基于公开数据自动生成，仅供研究参考，不构成投资建议。",
                  size=9, color=ON_DARK_SUB)
    set_notes(s, "收口：重申生意/公司/账/风险/价格五句结论。")

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


