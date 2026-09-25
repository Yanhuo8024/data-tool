#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
report_generator.py - 自动化报告生成（多模板）
"""
import base64, io, os, re
from pathlib import Path
from datetime import datetime
import pandas as pd
import numpy as np

try:
    import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
    HAS_MPL = True
except ImportError:
    HAS_MPL = False

try:
    from core_config import OUT_DIR  # 统一入口，避免多处定义分叉
except ImportError:
    OUT_DIR = Path(__file__).parent / "output"  # core_config 不可用时兜底
OUT_DIR.mkdir(exist_ok=True)


def _fig_to_b64(fig):
    buf = io.BytesIO(); fig.savefig(buf, format='png', dpi=100, bbox_inches='tight', facecolor='#fff'); buf.seek(0)
    return f'<img src="data:image/png;base64,{base64.b64encode(buf.read()).decode()}" style="max-width:100%;">'


def _make_plot(plot_fn):
    """通用绘图包装器：检查MPL → 执行plot_fn → 返回HTML"""
    if not HAS_MPL: return ""
    try:
        fig = plot_fn()
        if fig is None: return ""
        html = _fig_to_b64(fig) if isinstance(fig, plt.Figure) else str(fig)
        plt.close('all'); return html
    except Exception as e: plt.close('all'); return ""


def _plot_distributions(df, num_cols, max_cols=6):
    if not num_cols: return ""
    cols = num_cols[:max_cols]; n = len(cols)
    def _draw():
        rows = (n + 2) // 3; fig, axes = plt.subplots(rows, 3, figsize=(12, 3*rows))
        axes = axes.flatten() if hasattr(axes, 'flatten') else [axes]
        for i, col in enumerate(cols):
            df[col].dropna().hist(bins=30, ax=axes[i], color='#FF6B6B', alpha=0.8)
            axes[i].set_title(col, fontsize=10); axes[i].tick_params(labelsize=8)
        for j in range(i+1, len(axes)): axes[j].set_visible(False)
        plt.tight_layout(); return fig
    return _make_plot(_draw)


def _plot_correlation_heatmap(df, num_cols):
    if len(num_cols) < 2: return ""
    cols = num_cols[:10]; corr = df[cols].corr()
    def _draw():
        fig, ax = plt.subplots(figsize=(8, 6))
        im = ax.imshow(corr, cmap='RdBu_r', vmin=-1, vmax=1)
        ax.set_xticks(range(len(cols))); ax.set_yticks(range(len(cols)))
        ax.set_xticklabels(cols, rotation=45, ha='right', fontsize=8)
        ax.set_yticklabels(cols, fontsize=8); plt.colorbar(im, ax=ax)
        ax.set_title('Correlation Matrix', fontsize=12); plt.tight_layout(); return fig
    return _make_plot(_draw)


def _plot_boxplots(df, num_cols):
    if not num_cols: return ""
    cols = num_cols[:8]
    def _draw():
        fig, ax = plt.subplots(figsize=(10, 4)); df[cols].boxplot(ax=ax)
        ax.tick_params(labelsize=8); plt.xticks(rotation=30, ha='right'); plt.tight_layout(); return fig
    return _make_plot(_draw)


def _plot_time_series(df, date_col=None, value_col=None):
    if date_col is None:
        for c in df.columns:
            if 'date' in c.lower() or '日期' in c or 'time' in c.lower():
                try: df[c] = pd.to_datetime(df[c]); date_col = c; break
                except Exception: pass
    if date_col is None or date_col not in df.columns: return ""
    num_cols = df.select_dtypes(include=[np.number]).columns.tolist()
    if not num_cols: return ""
    value_col = value_col or num_cols[0]
    def _draw():
        fig, ax = plt.subplots(figsize=(12, 4))
        ax.plot(df[date_col], df[value_col], color='#FF6B6B', linewidth=1)
        ax.set_title(f'{value_col} over time', fontsize=12); ax.tick_params(labelsize=8); plt.tight_layout(); return fig
    return _make_plot(_draw)


def _plot_bar_top_n(df, col, n=10, title=None):
    if col not in df.columns: return ""
    top = df[col].value_counts().head(n)
    def _draw():
        fig, ax = plt.subplots(figsize=(8, 4)); top.plot(kind='barh', ax=ax, color='#FF6B6B')
        ax.set_title(title or f'Top {n} {col}', fontsize=11); ax.tick_params(labelsize=8); plt.tight_layout(); return fig
    return _make_plot(_draw)


def _stats_table_html(stats):
    rows = ""
    for k, v in stats.items():
        rows += f'<tr><td>{k}</td><td>{v}</td></tr>'
    return f'<table><tr><th>Metric</th><th>Value</th></tr>{rows}</table>'


def _base_html(title, content, header_color="#1a237e"):
    return f'''<!DOCTYPE html><html><head><meta charset="utf-8"><title>{title}</title>
<style>
    body{{font-family:'Microsoft YaHei',sans-serif;max-width:900px;margin:0 auto;padding:30px;color:#333;}}
    h1{{color:{header_color};border-bottom:3px solid {header_color};padding-bottom:10px;}}
    h2{{color:{header_color};border-left:4px solid #FF6B6B;padding-left:12px;margin-top:30px;}}
    table{{border-collapse:collapse;width:100%;margin:15px 0;font-size:13px;}}
    th,td{{border:1px solid #ddd;padding:8px;text-align:center;}}
    th{{background:{header_color};color:#fff;}}
    tr:nth-child(even){{background:#f5f5f5;}}
    .chart{{margin:20px 0;text-align:center;}}
    .kpi{{display:flex;gap:15px;margin:20px 0;}}
    .kpi-card{{flex:1;background:#e8eaf6;padding:15px;border-radius:8px;text-align:center;}}
    .kpi-value{{font-size:24px;font-weight:bold;color:#1565c0;}}
    .footer{{margin-top:40px;padding-top:15px;border-top:1px solid #ccc;color:#888;font-size:12px;text-align:center;}}
</style></head><body>{content}
<div class="footer">Generated by data.tool | {datetime.now().strftime('%Y-%m-%d')}</div></body></html>'''


def template_monthly_quant(df, file_name, stats, num_cols, cat_cols):
    today = datetime.now().strftime('%Y年%m月%d日')
    return_analysis = "".join([f"<tr><td>{col}</td><td>{df[col].mean():.4f}</td><td>{df[col].std():.4f}</td><td>{(df[col].std()/df[col].mean()*100 if df[col].mean() else 0):.2f}%</td></tr>" for col in num_cols[:3]])
    charts = _plot_time_series(df) + _plot_distributions(df, num_cols) + _plot_correlation_heatmap(df, num_cols) + _plot_boxplots(df, num_cols)
    content = f'<h1>A股月度量化分析报告</h1><p>Date: {today} | Source: {file_name}</p><div class="summary"><h3>Overview</h3><p>Rows: <strong>{stats["rows"]}</strong> | Num cols: <strong>{stats["num_cols"]}</strong> | Missing: <strong>{stats["missing_pct"]}</strong>%</p></div><h2>1. Descriptive Stats</h2>{_stats_table_html(stats)}<h2>2. Return Analysis</h2><table><tr><th>Metric</th><th>Mean</th><th>Std</th><th>CV</th></tr>{return_analysis}</table><h2>3. Charts</h2>{charts}'
    return _base_html("A股月度量化分析报告", content)


def template_weekly_sales(df, file_name, stats, num_cols, cat_cols):
    today = datetime.now().strftime('%Y年%m月%d日')
    sales = "".join([f"<tr><td>{col}</td><td>{df[col].sum():,.2f}</td><td>{df[col].mean():,.2f}</td><td>{df[col].max():,.2f}</td><td>{df[col].min():,.2f}</td></tr>" for col in num_cols[:5]])
    charts = (_plot_bar_top_n(df, cat_cols[0], 10, f'Top 10 {cat_cols[0]}') if cat_cols else "") + _plot_distributions(df, num_cols) + _plot_time_series(df)
    content = f'<h1>电商销售数据周报</h1><p>Date: {today} | Source: {file_name}</p><div class="kpi"><div class="kpi-card"><div class="kpi-value">{stats["rows"]}</div><div>Orders</div></div><div class="kpi-card"><div class="kpi-value">{stats["num_cols"]}</div><div>Dimensions</div></div><div class="kpi-card"><div class="kpi-value">{stats["missing_pct"]}%</div><div>Missing</div></div></div><h2>1. Sales Summary</h2><table><tr><th>Metric</th><th>Total</th><th>Mean</th><th>Max</th><th>Min</th></tr>{sales}</table><h2>2. Analysis</h2>{charts}'
    return _base_html("电商销售数据周报", content, "#e65100")


def template_data_exploration(df, file_name, stats, num_cols, cat_cols):
    today = datetime.now().strftime('%Y年%m月%d日')
    detail = "".join([f"<tr><td>{col}</td><td>{df[col].mean():.4f}</td><td>{df[col].median():.4f}</td><td>{df[col].std():.4f}</td><td>{df[col].min():.4f}</td><td>{df[col].max():.4f}</td><td>{df[col].skew():.4f}</td></tr>" for col in num_cols[:10] if len(df[col].dropna()) > 0])
    charts = _plot_distributions(df, num_cols) + _plot_correlation_heatmap(df, num_cols) + _plot_boxplots(df, num_cols) + (_plot_bar_top_n(df, cat_cols[0]) if cat_cols else "")
    content = f'<h1>数据探索性分析报告 (EDA)</h1><p>Date: {today} | Source: {file_name}</p><h2>1. Data Quality</h2>{_stats_table_html(stats)}<h2>2. Descriptive Stats</h2><table><tr><th>Col</th><th>Mean</th><th>Median</th><th>Std</th><th>Min</th><th>Max</th><th>Skew</th></tr>{detail}</table><h2>3. Charts</h2>{charts}'
    return _base_html("数据探索性分析报告", content, "#1b5e20")


def template_financial_statement(df, file_name, stats, num_cols, cat_cols):
    today = datetime.now().strftime('%Y年%m月%d日')
    fin = "".join([f"<tr><td>{col}</td><td>{df[col].sum():,.2f}</td><td>{df[col].mean():,.2f}</td><td>{df[col].std():,.2f}</td><td>{getattr(df[col].pct_change(), 'mean', lambda: 0)()*100:.2f}%</td></tr>" for col in num_cols[:8] if len(df[col].dropna()) > 0])
    charts = _plot_time_series(df) + _plot_distributions(df, num_cols) + _plot_boxplots(df, num_cols)
    content = f'<h1>财务分析报告</h1><p>Date: {today} | Source: {file_name}</p><h2>1. Overview</h2>{_stats_table_html(stats)}<h2>2. Financial Metrics</h2><table><tr><th>Metric</th><th>Total</th><th>Mean</th><th>Std</th><th>Change</th></tr>{fin}</table><h2>3. Charts</h2>{charts}'
    return _base_html("财务分析报告", content, "#4a148c")


TEMPLATES = {
    "monthly_quant": {"name": "A股月度量化分析报告", "description": "股票K线月度量化分析", "func": template_monthly_quant},
    "weekly_sales": {"name": "电商销售数据周报", "description": "电商销售周报", "func": template_weekly_sales},
    "data_exploration": {"name": "数据探索性分析报告", "description": "通用EDA报告", "func": template_data_exploration},
    "financial_statement": {"name": "财务分析报告", "description": "财务数据分析", "func": template_financial_statement},
}


def list_templates():
    return {k: {"name": v["name"], "description": v["description"]} for k, v in TEMPLATES.items()}


def generate_report(file_path, template="data_exploration", output_format="html"):
    try:
        p = Path(file_path)
        if not p.exists(): return None, {"error": f"File not found: {file_path}"}
        if p.suffix.lower() == '.csv': df = pd.read_csv(p, encoding='utf-8-sig')
        elif p.suffix.lower() in ('.xlsx', '.xls'): df = pd.read_excel(p)
        else: return None, {"error": "Unsupported format"}
        df = df.copy()
        num_cols = df.select_dtypes(include=[np.number]).columns.tolist()
        cat_cols = df.select_dtypes(include=['object', 'category']).columns.tolist()
        stats = {'file': p.name, 'rows': len(df), 'cols': len(df.columns),
                 'num_cols': len(num_cols), 'cat_cols': len(cat_cols),
                 'missing_cells': int(df.isna().sum().sum()),
                 'missing_pct': round(df.isna().sum().sum() / max(len(df)*len(df.columns), 1) * 100, 1),
                 'duplicate_rows': int(df.duplicated().sum())}
        tmpl = TEMPLATES.get(template)
        if not tmpl: return None, {"error": f"Unknown template: {template}"}
        html_content = tmpl["func"](df, p.name, stats, num_cols, cat_cols)
        if output_format == "docx":
            try:
                from docx import Document
                doc = Document(); doc.add_heading(tmpl["name"], 0)
                text = re.sub(r'<[^>]+>', ' ', html_content); text = re.sub(r'\s+', ' ', text).strip()
                doc.add_paragraph(text[:5000])
                buf = io.BytesIO(); doc.save(buf); buf.seek(0)
                return buf.read(), {"format": "docx", "template": template}
            except ImportError: pass
        return html_content.encode('utf-8'), {"format": "html", "template": template}
    except Exception as e: return None, {"error": str(e)}


def save_report(file_path, output_dir=None, template="data_exploration"):
    output_dir = output_dir or OUT_DIR; os.makedirs(output_dir, exist_ok=True)
    content, meta = generate_report(file_path, template, "html")
    if content is None: return None, meta
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"report_{meta['template']}_{ts}.html"
    filepath = os.path.join(output_dir, filename)
    with open(filepath, 'wb') as f: f.write(content)
    return filepath, meta
